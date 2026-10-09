#!/usr/bin/env python3
"""Checks that a submitted file is read completely -- or reported as incomplete.

Builds real .docx, .pdf and .txt files in memory and runs them through the same
extractor the grading job uses. No network, no API key:

    python3 scripts/test_extraction.py

Every fixture is built in memory, so this needs no files on disk.

The case that motivates all of this: a student who answers inside a Word table
used to extract as their name alone, get graded, and score zero with a confident
justification. Nothing anywhere said extraction had failed. Every check here
exists to make sure that class of silent wrong grade cannot happen again.
"""
import io
import json
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402
from docx.shared import Inches  # noqa: E402

from app import extract  # noqa: E402

FAILS = []

ANSWER_1 = "A BST keeps left subtree values smaller than the node. Average case O(log n)."
ANSWER_2 = "Naive recursion is O(2^n) because it recomputes overlapping subproblems."


def check(label, condition):
    print(("PASS  " if condition else "FAIL  ") + label)
    if not condition:
        FAILS.append(label)


def docx_bytes(build):
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def read(filename, content):
    return extract.extract_document(filename, content)


# --- .docx --------------------------------------------------------------------

def test_docx_tables():
    print("\n-- answers inside a Word table (the silent-zero bug) --")

    def build(d):
        d.add_paragraph("Name: Jordan Lee")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Q1"
        table.cell(0, 1).text = ANSWER_1
        table.cell(1, 0).text = "Q2"
        table.cell(1, 1).text = ANSWER_2

    result = read("table.docx", docx_bytes(build))
    check("a table answer is recovered", ANSWER_1 in result.text and ANSWER_2 in result.text)
    check("the question labels come with it", "Q1" in result.text and "Q2" in result.text)
    check("the table was counted", result.stats["tables"] == 1)
    # A short answer laid out in a table is a short answer, not a read failure.
    check("a fully-read table is not flagged incomplete", not result.looks_incomplete)

    def build_nested(d):
        outer = d.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = ANSWER_1

    nested = read("nested.docx", docx_bytes(build_nested))
    check("a table nested inside a table is recovered", ANSWER_1 in nested.text)


def test_docx_reading_order():
    print("\n-- document order is preserved --")

    def build(d):
        d.add_paragraph("Q1 answer follows.")
        t1 = d.add_table(rows=1, cols=1)
        t1.cell(0, 0).text = ANSWER_1
        d.add_paragraph("Q2 answer follows.")
        t2 = d.add_table(rows=1, cols=1)
        t2.cell(0, 0).text = ANSWER_2

    text = read("order.docx", docx_bytes(build)).text
    positions = [text.index(s) for s in
                 ("Q1 answer follows", ANSWER_1, "Q2 answer follows", ANSWER_2)]
    check("paragraphs and tables interleave in document order",
          positions == sorted(positions))


def test_docx_other_containers():
    print("\n-- text boxes, equations, hyperlinks, list items --")

    def build_list(d):
        d.add_paragraph(ANSWER_1, style="List Bullet")
        d.add_paragraph(ANSWER_2, style="List Number")

    listed = read("list.docx", docx_bytes(build_list))
    check("bulleted and numbered answers are read",
          ANSWER_1 in listed.text and ANSWER_2 in listed.text)

    # A hyperlink's label lives in a w:hyperlink run, which paragraph.text drops.
    def build_link(d):
        p = d.add_paragraph("See ")
        run = p.add_run()
        link = run._element.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink", {})
        inner_run = run._element.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r", {})
        text_el = run._element.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t", {})
        text_el.text = "merge sort is O(n log n)"
        inner_run.append(text_el)
        link.append(inner_run)
        p._p.append(link)

    linked = read("link.docx", docx_bytes(build_link))
    check("hyperlink text is read", "merge sort is O(n log n)" in linked.text)

    # An equation-editor formula: <m:oMath><m:r><m:t>O(n log n)</m:t>
    def build_math(d):
        p = d.add_paragraph("Complexity: ")
        omath = p._p.makeelement(
            "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath", {})
        m_r = p._p.makeelement(
            "{http://schemas.openxmlformats.org/officeDocument/2006/math}r", {})
        m_t = p._p.makeelement(
            "{http://schemas.openxmlformats.org/officeDocument/2006/math}t", {})
        m_t.text = "O(n log n)"
        m_r.append(m_t)
        omath.append(m_r)
        p._p.append(omath)

    math = read("math.docx", docx_bytes(build_math))
    check("equation-editor text is read", "O(n log n)" in math.text)
    check("equations are counted", math.stats["equations"] == 1)
    check("a note warns that notation may be imprecise",
          any("equation" in n for n in math.notes))


def test_docx_incomplete_detection():
    print("\n-- an unreadable document is refused, not graded --")

    # A picture of the answers plus a name: the exact shape of a submission that
    # would otherwise score a confident zero.
    png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08"
           b"\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00"
           b"\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")

    def build(d):
        d.add_paragraph("Name: Jordan Lee")
        d.add_picture(io.BytesIO(png), width=Inches(1))

    result = read("scan.docx", docx_bytes(build))
    check("the image is counted", result.stats["images"] == 1)
    check("an image-only document is flagged incomplete", result.looks_incomplete)
    check("a note names the images", any("image" in n for n in result.notes))

    # The other side of that judgment: a genuinely short answer, no images, must
    # still be graded so it can score what it deserves.
    short = read("short.docx", docx_bytes(lambda d: d.add_paragraph("Q1: A tree.")))
    check("a genuinely short answer is NOT flagged incomplete", not short.looks_incomplete)
    check("the short answer's text is intact", "A tree" in short.text)

    def build_empty(d):
        d.add_paragraph("")

    try:
        read("empty.docx", docx_bytes(build_empty))
        check("an empty document raises", False)
    except extract.ExtractionError:
        check("an empty document raises", True)

    try:
        read("broken.docx", b"this is not a docx at all")
        check("a corrupt docx raises a readable error", False)
    except extract.ExtractionError as exc:
        check("a corrupt docx raises a readable error", "could not be opened" in str(exc))


def test_docx_header_fallback():
    print("\n-- header/footer used only as a fallback --")

    def build_normal(d):
        d.add_paragraph(ANSWER_1)
        d.add_paragraph(ANSWER_2)
        d.sections[0].header.paragraphs[0].text = "Jordan Lee - CSCI 101"

    normal = read("hdr.docx", docx_bytes(build_normal))
    check("a normal document does not pull in header noise",
          "CSCI 101" not in normal.text)

    def build_body_in_header(d):
        d.add_paragraph("Name")
        d.sections[0].header.paragraphs[0].text = ANSWER_1

    fallback = read("hdr2.docx", docx_bytes(build_body_in_header))
    check("a nearly empty body falls back to the header", ANSWER_1 in fallback.text)
    check("the fallback is noted", any("header" in n.lower() for n in fallback.notes))


# --- .pdf ---------------------------------------------------------------------

TINY_PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08"
            b"\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00"
            b"\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")


def pdf_bytes(lines, with_image=False):
    """A minimal single-page PDF with a real text layer, optionally an image too."""
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdf_canvas

    buffer = io.BytesIO()
    c = pdf_canvas.Canvas(buffer, pagesize=LETTER)
    y = 720
    for line in lines:
        c.drawString(72, y, line)
        y -= 18
    if with_image:
        c.drawImage(ImageReader(io.BytesIO(TINY_PNG)), 72, 500, width=72, height=72, mask="auto")
    c.save()
    return buffer.getvalue()


def test_pdf():
    print("\n-- .pdf --")
    try:
        import reportlab  # noqa: F401
    except ImportError:
        print("SKIP  reportlab not installed; install it to test PDF extraction")
        return

    result = read("answers.pdf", pdf_bytes([f"Q1: {ANSWER_1}", f"Q2: {ANSWER_2}"]))
    check("pdf text is extracted", ANSWER_1 in result.text and ANSWER_2 in result.text)
    check("the page count is recorded", result.stats["pages"] == 1)
    check("a normal pdf is not flagged incomplete", not result.looks_incomplete)

    multi = read("multi.pdf", pdf_bytes([f"Q{i}: line {i}" for i in range(1, 30)]))
    check("every line of a longer pdf is read",
          all(f"line {i}" in multi.text for i in range(1, 30)))

    # Typed answers plus a diagram: the text grades fine, but the diagram was
    # skipped and the faculty member has to be told so.
    mixed = read("mixed.pdf", pdf_bytes([f"Q{i}: {ANSWER_1}" for i in range(1, 12)],
                                        with_image=True))
    check("an image alongside readable text is counted", mixed.stats["images"] >= 1)
    check("an image alongside readable text is reported",
          any("cannot read" in n for n in mixed.notes))
    check("text plus an image still grades (not refused)", not mixed.looks_incomplete)
    check("the text itself is intact", ANSWER_1 in mixed.text)

    # The other side: an image and almost no text is a scan, and is refused.
    scan = read("scanned.pdf", pdf_bytes(["p1"], with_image=True))
    check("an image with almost no text is refused as a scan", scan.looks_incomplete)

    plain_sparse = read("sparse.pdf", pdf_bytes(["hi"]))
    check("sparse text with no images is noted but still graded",
          not plain_sparse.looks_incomplete
          and any("very little" in n.lower() for n in plain_sparse.notes))

    try:
        read("broken.pdf", b"%PDF-1.4 truncated garbage")
        check("a corrupt pdf raises a readable error", False)
    except extract.ExtractionError as exc:
        check("a corrupt pdf raises a readable error", "could not be read" in str(exc))


# --- .txt ---------------------------------------------------------------------

def test_txt_encodings():
    print("\n-- .txt encodings --")

    plain = read("a.txt", ANSWER_1.encode("utf-8"))
    check("utf-8 is read", plain.text == ANSWER_1)
    check("the encoding is recorded", plain.stats["encoding"] == "utf-8")

    bom = read("b.txt", b"\xef\xbb\xbf" + ANSWER_1.encode("utf-8"))
    check("a utf-8 BOM is stripped, not read as a character",
          bom.text == ANSWER_1 and "﻿" not in bom.text)

    utf16 = read("c.txt", ANSWER_1.encode("utf-16"))
    check("utf-16 (Word's 'save as plain text') is read", utf16.text == ANSWER_1)

    # cp1252 smart quotes -- what Word produces, and not valid UTF-8.
    # cp1252 bytes directly: 0x93/0x94 are smart quotes, 0x96 an en dash.
    windows = b"Q1: the tree is \x93degenerate\x94 \x96 O(n)."
    cp = read("d.txt", windows)
    check("a cp1252 file is decoded rather than mangled",
          "degenerate" in cp.text and "�" not in cp.text)
    check("the non-utf-8 encoding is noted", any("UTF-8" in n for n in cp.notes))

    try:
        read("e.txt", b"")
        check("an empty file raises", False)
    except extract.ExtractionError:
        check("an empty file raises", True)


# --- refusals and text entry ---------------------------------------------------

def notebook(cells, **top):
    payload = {"nbformat": 4, "cells": cells}
    payload.update(top)
    return json.dumps(payload).encode()


def code_cell(source, count=1, outputs=()):
    return {"cell_type": "code", "execution_count": count,
            "source": source, "outputs": list(outputs)}


def test_notebook():
    print("\n-- .ipynb: cells, outputs, and whether the code was ever run --")

    result = read("hw.ipynb", notebook([
        {"cell_type": "markdown", "source": ["# Step 1\n", "Read the data."]},
        code_cell(["import pandas as pd\n", "df = pd.read_csv('f.csv')\n", "df.shape"],
                  count=1,
                  outputs=[{"output_type": "execute_result",
                            "data": {"text/plain": ["(7752, 25)"]}}]),
    ]))
    check("markdown cells are read", "Read the data." in result.text)
    check("code is kept verbatim so it can be quoted as evidence",
          "df = pd.read_csv('f.csv')" in result.text)
    check("the output of a cell is included", "(7752, 25)" in result.text)
    check("an executed cell says so", "ran as In[1]" in result.text)
    check("cells keep notebook order",
          result.text.index("# Step 1") < result.text.index("import pandas"))

    # The distinction the whole reader exists for: written but never run.
    result = read("hw.ipynb", notebook([
        code_cell(["df.drop_duplicates(inplace=True)"], count=None, outputs=[]),
    ]))
    check("a cell that was never run is labelled", "NEVER RUN" in result.text)
    check("a cell that was never run is reported to faculty",
          any("never" in note for note in result.notes))
    check("the code of an unrun cell is still read",
          "df.drop_duplicates(inplace=True)" in result.text)

    result = read("hw.ipynb", notebook([
        code_cell(["df.drop_duplictes()"], count=4, outputs=[{
            "output_type": "error", "ename": "AttributeError", "evalue": "nope",
            "traceback": ["\u001b[0;31mAttributeError\u001b[0m: no attribute 'drop_duplictes'"]}]),
    ]))
    check("a cell that raised is labelled", "--- ERROR ---" in result.text)
    check("the traceback is readable", "no attribute 'drop_duplictes'" in result.text)
    check("terminal colour codes are stripped", "\u001b[" not in result.text)
    check("an error is reported to faculty",
          any("error" in note for note in result.notes))

    # A plot is the deliverable for several steps. We cannot see it, but the
    # fact that one was produced is real evidence and must not be silently lost.
    result = read("hw.ipynb", notebook([
        code_cell(["sns.heatmap(df.corr())"], count=2, outputs=[{
            "output_type": "display_data",
            "data": {"image/png": "iVBORw0KGgo=", "text/plain": ["<Axes: >"]}}]),
    ]))
    check("a figure in the output is recorded", "a figure was produced" in result.text)
    check("the unreadable figure is disclosed to faculty",
          any("could not be read" in note for note in result.notes))
    check("a notebook with figures is not treated as a failed read",
          not result.looks_incomplete)

    # One cell can print seven thousand rows; the next question still has to fit.
    huge = read("hw.ipynb", notebook([
        code_cell(["df"], count=1, outputs=[{"output_type": "stream", "name": "stdout",
                                             "text": ["row %d\n" % i for i in range(500)]}]),
    ]))
    check("a runaway output is truncated", "[output truncated]" in huge.text)
    check("truncation keeps the beginning of the output", "row 0" in huge.text)
    check("truncation is bounded", len(huge.text) < 4000)

    # Jupyter leaves an empty cell at the bottom of almost every notebook. A
    # real submission tripped this: a complete paper carrying a false warning.
    tidy = read("hw.ipynb", notebook([
        code_cell(["df.head()"], count=1, outputs=[{"output_type": "stream",
                                                    "name": "stdout", "text": ["ok"]}]),
        {"cell_type": "code", "execution_count": None, "source": [], "outputs": []},
    ]))
    check("an empty trailing cell is not counted as code", tidy.stats["code_cells"] == 1)
    check("an empty trailing cell raises no warning", tidy.notes == [])

    check("a notebook with no code cells is reported, not refused",
          any("no code cells" in note for note in
              read("hw.ipynb", notebook([{"cell_type": "markdown",
                                          "source": ["just prose"]}])).notes))

    for payload, expected in ((b"not json at all", "not valid notebook JSON"),
                              (notebook([]), "no cells")):
        try:
            read("hw.ipynb", payload)
            check(f"a broken notebook is refused ({expected})", False)
        except extract.ExtractionError as exc:
            check(f"a broken notebook is refused ({expected})", expected in str(exc))

    check("a plain .py script is read as text",
          "import pandas" in read("hw.py", b"import pandas as pd\n").text)


def _png(text="handwritten answer"):
    from PIL import Image, ImageDraw
    image = Image.new("RGB", (900, 300), "white")
    ImageDraw.Draw(image).text((30, 30), text, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _scanned_pdf(pages=2):
    """A PDF whose pages are pictures of words -- no text layer at all."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as rl

    picture = io.BytesIO(_png())
    out = io.BytesIO()
    pdf = rl.Canvas(out, pagesize=(612, 792))
    for _ in range(pages):
        picture.seek(0)
        pdf.drawImage(ImageReader(picture), 40, 400, width=500, height=250)
        pdf.showPage()
    pdf.save()
    return out.getvalue()


def _with_backend(name, transcript):
    """Reload extract/ocr with a backend selected and the reader stubbed."""
    import importlib
    from unittest import mock
    import app.extract
    import app.ocr

    os.environ["OCR_BACKEND"] = name
    importlib.reload(app.ocr)
    importlib.reload(app.extract)
    reader = lambda image, index, context, **kwargs: app.ocr.PageText(index, transcript)  # noqa: E731
    return app.ocr, app.extract, mock.patch.object(app.ocr, "_read_page_vision", side_effect=reader)


def _restore():
    import importlib
    import app.extract
    import app.ocr
    os.environ.pop("OCR_BACKEND", None)
    importlib.reload(app.ocr)
    importlib.reload(app.extract)


TRANSCRIPT = ("Q1. A BST keps smaller values on the left.\n"
              "Worst case is O(n) when [?unbalanced].\n"
              "[struck: average is O(n^2)]\n"
              "[diagram: a tree with 7 nodes]\n"
              "[illegible]")


def test_handwriting_off_by_default():
    print("\n-- with no reader configured, images are refused, not graded --")
    from app import extract as ex

    check("image formats are not accepted by default",
          "png" not in ex.supported_extensions())
    try:
        ex.extract_document("page.png", _png())
        check("a photograph is refused when reading is off", False)
    except ex.ExtractionError as exc:
        # The refusal must say the feature is OFF, not that it does not exist:
        # a faculty member told "images cannot be read" chases the student for
        # a different file when the fix is a server setting.
        check("a photograph is refused when reading is off",
              "turned off on this server" in str(exc) and "OCR_BACKEND" in str(exc))
    try:
        ex.extract_document("scan.pdf", _scanned_pdf())
        check("a scanned PDF is refused when reading is off", False)
    except ex.ExtractionError as exc:
        check("a scanned PDF is refused when reading is off",
              "scan or photographs" in str(exc) and "turned off on this server" in str(exc))


def test_handwriting_transcription():
    print("\n-- handwriting and scans, once a reader is configured --")
    ocr, ex, patch = _with_backend("vision", TRANSCRIPT)
    try:
        check("image formats become acceptable", "png" in ex.supported_extensions())

        with patch:
            result = ex.extract_document("page.png", _png(), "DSA Midterm")

        check("the transcript reaches the grader", "A BST keps smaller values" in result.text)
        check("the page is labelled as a machine transcription",
              "machine transcription" in result.text)
        check("the submission is marked as transcribed, not read",
              result.stats["from_ocr"] is True)
        check("a transcribed page is not treated as a failed read",
              not result.looks_incomplete)
        check("faculty are told the grade rests on a reading of the writing",
              any("check it against the original" in n for n in result.notes))
        check("uncertainty markers are counted for the reviewer",
              result.stats["ocr_markers"] == 2)

        # The failure that no downstream check could ever catch.
        check("the student's own misspelling survives transcription",
              "keps" in result.text and "keeps smaller" not in result.text)

        with patch:
            scanned = ex.extract_document("scan.pdf", _scanned_pdf(2), "DSA")
        check("a scanned PDF is rasterised and transcribed page by page",
              scanned.stats["pages"] == 2 and scanned.stats["from_ocr"] is True)
        check("a transcribed scan is graded rather than refused",
              not scanned.looks_incomplete and scanned.chars > 0)
        check("the scan says why it needed transcribing",
              any("no text layer" in n for n in scanned.notes))

        # Colour PNG pages made a real 11-page submission 19 MB to send. Ink on
        # paper has no colour, so this is cost with nothing bought.
        rendered = ex._render_pdf_pages(_scanned_pdf(1))
        check("rendered pages are grayscale JPEG, not colour PNG",
              rendered[0][:3] == b"\xff\xd8\xff")
        check("the reader declares the format it is actually sending",
              ocr._media_type(rendered[0]) == "image/jpeg"
              and ocr._media_type(_png()) == "image/png")

        # A reader that fails must not hand back an empty submission, which
        # would be graded as though the student wrote nothing.
        from unittest import mock
        with mock.patch.object(ocr, "_read_page_vision",
                               side_effect=ocr.OCRError("reader unavailable")):
            try:
                ex.extract_document("page.png", _png())
                check("a failed read refuses instead of grading a blank", False)
            except ex.ExtractionError as exc:
                check("a failed read refuses instead of grading a blank",
                      "reader unavailable" in str(exc))

        with mock.patch.object(ocr, "_read_page_vision",
                               side_effect=lambda i, n, c, **kwargs: ocr.PageText(n, "")):
            try:
                ex.extract_document("page.png", _png())
                check("an empty transcript refuses instead of grading a blank", False)
            except ex.ExtractionError:
                check("an empty transcript refuses instead of grading a blank", True)
    finally:
        _restore()


def test_vision_gateway_shape():
    print("\n-- the request shape and model, as measured against IU's gateway --")
    import importlib
    import app.ocr

    os.environ["OCR_BACKEND"] = "vision"
    os.environ.pop("OCR_VISION_MODEL", None)
    os.environ.pop("OCR_TRY_ANTHROPIC_SHAPE", None)
    importlib.reload(app.ocr)
    o = app.ocr
    try:
        check("the default transcriber is a model probed to actually see images",
              o.VISION_MODEL == "gemma-4-31B-it")
        # The gateway answers the Anthropic block with a 500, so trying it
        # spends a failed request on every page of every submission.
        check("only the shape this gateway accepts is sent by default",
              len(o._vision_payloads(b"\x89PNG", 1, "ctx")) == 1)
        check("the sent shape is the OpenAI one",
              o._vision_payloads(b"\x89PNG", 1, "ctx")[0][1]["type"] == "image_url")

        os.environ["OCR_TRY_ANTHROPIC_SHAPE"] = "1"
        importlib.reload(app.ocr)
        check("the other shape is still available for a different gateway",
              len(app.ocr._vision_payloads(b"\x89PNG", 1, "ctx")) == 2)
        os.environ.pop("OCR_TRY_ANTHROPIC_SHAPE", None)

        # A document-conversion model answers in layout markup.
        doctags = ("<loc_0><loc_0><loc_500><loc_500><_unknown>"
                   "B <- empty array </line_code>")
        check("layout markup is stripped before the grader sees it",
              o.strip_doctags(doctags) == "B <- empty array")
        check("ordinary transcription is left alone",
              o.strip_doctags("A BST keeps values [?left]") == "A BST keeps values [?left]")
    finally:
        _restore()


def test_an_empty_reply_is_retried_not_fatal():
    """The failure that cost a real submission: "the model returned nothing".

    It happened on page 2 of 11. One page refusing refuses the whole
    submission -- deliberately, so that a partial read is never graded as a
    whole one -- but this page was never asked a second time. A timeout got
    three attempts; an empty reply and a 503 got one each, though neither is
    any more a verdict on the page than a timeout is.

    The other half of it: this gateway answers in reasoning_content for some
    models and leaves content empty. app/llm_client.py has read both fields
    since the beginning. The transcriber read only content, so a page the
    model HAD transcribed was reported as nothing at all.
    """
    import importlib
    from unittest import mock

    import app.ocr

    print("\n-- a page with no answer is asked again, not written off --")
    os.environ["OCR_BACKEND"] = "vision"
    os.environ["REALLM_URL"] = "https://example.invalid/v1/chat/completions"
    os.environ["REALLM_API_KEY"] = "test-only"
    os.environ["OCR_BACKOFF_SECONDS"] = "0"
    importlib.reload(app.ocr)
    o = app.ocr

    class Reply:
        def __init__(self, status=200, content="", finish="stop", reasoning=None):
            self.status_code = status
            message = {"content": content}
            if reasoning is not None:
                message["reasoning_content"] = reasoning
            self._body = {"choices": [{"message": message, "finish_reason": finish}]}
            self.text = "server error"

        def json(self):
            return self._body

    try:
        # The gateway answered in the other field.
        with mock.patch.object(
                o.requests, "post",
                side_effect=lambda *a, **k: Reply(content="", reasoning="A BST keeps smaller values left.")):
            page = o._read_page_vision(b"\x89PNG", 1, "ctx")
        check("a reply carried in reasoning_content is read, not discarded",
              "smaller values" in page.text)

        # Empty once, then fine: exactly the real failure, retried.
        calls = []

        def flaky(*_a, **_k):
            calls.append(1)
            if len(calls) == 1:
                return Reply(content="")
            return Reply(content="Problem 2. Merge sort is O(n lg n).")

        with mock.patch.object(o.requests, "post", side_effect=flaky):
            page = o._read_with_retries(o._read_page_vision, b"\x89PNG", 2, "ctx")
        check("an empty reply is asked again", len(calls) == 2)
        check("the second answer is used", "Merge sort" in page.text)

        # A 503 is the gateway's problem, not the page's.
        calls.clear()

        def unavailable(*_a, **_k):
            calls.append(1)
            if len(calls) < 3:
                return Reply(status=503)
            return Reply(content="Problem 3. The recurrence solves to O(n).")

        with mock.patch.object(o.requests, "post", side_effect=unavailable):
            page = o._read_with_retries(o._read_page_vision, b"\x89PNG", 3, "ctx")
        check("a 503 is retried rather than refusing the submission", len(calls) == 3)
        check("the page is read once the gateway recovers", "O(n)" in page.text)

        # But a request that is simply wrong must not be repeated.
        calls.clear()
        with mock.patch.object(o.requests, "post",
                               side_effect=lambda *a, **k: (calls.append(1), Reply(status=400))[1]):
            try:
                o._read_with_retries(o._read_page_vision, b"\x89PNG", 4, "ctx")
            except o.OCRError:
                pass
        check("a 400 is not retried, because asking again cannot fix it",
              len(calls) == 1)

        # Still gives up eventually, rather than retrying for ever.
        calls.clear()
        with mock.patch.object(o.requests, "post",
                               side_effect=lambda *a, **k: (calls.append(1), Reply(content=""))[1]):
            try:
                o._read_with_retries(o._read_page_vision, b"\x89PNG", 5, "ctx")
                check("an always-empty page still fails in the end", False)
            except o.OCRError as exc:
                message = str(exc)
                check("an always-empty page still fails in the end", True)
                check("and says how many attempts were made",
                      str(o.MAX_ATTEMPTS) in message)
                check("the error names the finish_reason, so the next failure "
                      "is diagnosable", "finish_reason" in message)
        check("it stops at the attempt limit", len(calls) == o.MAX_ATTEMPTS)
    finally:
        for key in ("REALLM_URL", "REALLM_API_KEY", "OCR_BACKOFF_SECONDS"):
            os.environ.pop(key, None)
        _restore()


def test_runaway_generation_is_bounded_and_detected():
    """A page that sends the model into a loop must not cost 240s, or be graded.

    Nine pages of a real submission answered in 8-11s; two ran to the 240s
    timeout, and those two were 83% of the time to read it. A page is about
    450 output tokens, so a 240s reply at any plausible rate means thousands
    of tokens -- the model did not stop. Nothing capped it (no max_tokens) and
    nothing noticed (nothing read finish_reason), so a reply cut off at the
    server's own limit was stored as if it were the whole page.
    """
    import importlib
    import json as jsonlib
    from unittest import mock

    import app.ocr

    print("\n-- a transcriber that will not stop is capped and caught --")
    os.environ["OCR_BACKEND"] = "vision"
    os.environ["REALLM_URL"] = "https://example.invalid/v1/chat/completions"
    os.environ["REALLM_API_KEY"] = "test-only"
    importlib.reload(app.ocr)
    o = app.ocr

    class Reply:
        status_code = 200

        def __init__(self, text, finish="stop"):
            self._body = {"choices": [{"message": {"content": text},
                                       "finish_reason": finish}]}

        def json(self):
            return self._body

    try:
        sent = {}

        def capture(url, **kwargs):
            sent.update(kwargs.get("json") or {})
            return Reply("A BST keeps smaller values on the left.")

        with mock.patch.object(o.requests, "post", side_effect=capture):
            page = o._read_page_vision(b"\x89PNG", 1, "ctx")
        check("a token ceiling is sent with every page", "max_tokens" in sent)
        check(f"the ceiling is generous, not tight ({sent.get('max_tokens')})",
              1000 <= sent.get("max_tokens", 0) <= 8000)
        check("a normal page still comes back", "smaller values" in page.text)
        check("the request is still deterministic", sent.get("temperature") == 0)

        # Cut off at the limit, but not repetitive: a genuinely long page.
        # It must be refused, never stored as a complete reading.
        # Varied prose, not a repeated phrase: a page that is genuinely long is
        # a different thing from a page that is looping, and the two must not
        # be answered the same way.
        long_page = "\n".join(
            f"Step {i}: substituting the bound from step {i - 1} gives "
            f"T(n) <= {i}n lg n + {i * 3}n, which is still O(n lg n)."
            for i in range(1, 40))
        budgets = []

        def truncating(*_a, **kwargs):
            budgets.append((kwargs.get("json") or {}).get("max_tokens"))
            return Reply(long_page, "length")

        with mock.patch.object(o.requests, "post", side_effect=truncating):
            try:
                o._read_page_vision(b"\x89PNG", 2, "ctx")
                check("a truncated page is not accepted as complete", False)
            except o.TransientOCRError:
                check("a truncated page is not accepted as complete", True)

            # It is asked again with more room before anything is refused. A
            # fixed ceiling refused a real submission whose page needed about
            # 250 tokens and had produced 2000 -- ample room, and refusing was
            # still the wrong answer.
            budgets.clear()
            try:
                o._read_with_retries(o._read_page_vision, b"\x89PNG", 2, "ctx")
                check("a page that never finishes is refused in the end", False)
            except o.OCRError as exc:
                check("a page that never finishes is refused in the end", True)
                check("the refusal says what actually happened, not 'raise the limit'",
                      "without finishing" in str(exc))
        check(f"each attempt gets more room than the last ({budgets})",
              len(budgets) == o.MAX_ATTEMPTS and budgets[1] > budgets[0])
        check("the widening stops at a ceiling",
              max(budgets) <= o.MAX_OUTPUT_TOKENS_CEILING)

        # A repetition loop is a different failure and says so, because the
        # answer is a better scan, not a higher token limit.
        loop = "the the the the the the the the the the \n" * 30
        with mock.patch.object(o.requests, "post",
                               side_effect=lambda *a, **k: Reply(loop, "length")):
            try:
                o._read_page_vision(b"\x89PNG", 3, "ctx")
                check("a looping page is refused", False)
            except o.OCRError as exc:
                check("a looping page is refused", True)
                check("a loop is reported as a loop, not as a length limit",
                      "loop" in str(exc).lower())

        # The first version of the detector failed exactly here: thirty-nine
        # varied lines that happened to END the same way were counted as a
        # loop, because it asked how often the tail recurred anywhere rather
        # than whether it recurred back to back.
        check("varied prose with repeated line endings is not a loop",
              not o.looks_degenerate(long_page))
        # A real answer sheet is not periodic, and is not flagged.
        answers = "True False False True True False True False False False " \
                  "True True False True False False True True True False".split()
        check("a page of one-word answers is not a loop",
              not o.looks_degenerate("\n".join(f"{i}. {a}" for i, a in
                                               enumerate(answers * 5, start=1))))
        # A PERFECTLY cyclic page is indistinguishable from a loop by text
        # alone, and does read as one. That costs the wording of a refusal and
        # nothing else -- both branches refuse the page either way -- so it is
        # recorded here as known behaviour rather than papered over.
        check("a perfectly periodic page reads as a loop, which only changes the advice",
              o.looks_degenerate("\n".join(["True", "False"] * 60)))
        # A reply that ended by itself is trusted even if it looks repetitive:
        # the model decided it was done, and the alternative is refusing a
        # page of legitimately repetitive work.
        repetitive_but_finished = "\n".join(["True"] * 100)
        with mock.patch.object(
                o.requests, "post",
                side_effect=lambda *a, **k: Reply(repetitive_but_finished, "stop")):
            page = o._read_page_vision(b"\x89PNG", 4, "ctx")
        check("a repetitive page that ENDED is still accepted",
              page.text.startswith("True"))

        check("a real transcript is not mistaken for a loop",
              not o.looks_degenerate(
                  "Problem 1. We sort the array first, which costs O(n lg n).\n"
                  "Then a single pass finds the pair, so the total is O(n lg n).\n"
                  "[diagram: two pointers moving inward]\n"
                  "Problem 2. The recurrence is T(n) = 2T(n/2) + n, so by the\n"
                  "master theorem the answer is O(n lg n) again. I checked this\n"
                  "against the table in the notes and it matches case 2."))
        check("a short repeated answer is not mistaken for a loop",
              not o.looks_degenerate("True\nTrue\nFalse\nTrue"))
        check("an end-to-end repetition with no newlines is caught",
              o.looks_degenerate("the quick brown fox jumps over " * 40))
    finally:
        os.environ.pop("REALLM_URL", None)
        os.environ.pop("REALLM_API_KEY", None)
        _restore()


def test_pages_are_transcribed_in_parallel():
    """Serial transcription is hostage to its slowest page.

    Measured on a real submission: two pages of eleven took over 240s while the
    rest were quick, and the whole thing ran about ten minutes. In parallel the
    cost is roughly the slowest page rather than the sum.
    """
    import importlib
    import time
    from unittest import mock
    import app.ocr

    print("\n-- pages are read concurrently, in page order --")
    os.environ["OCR_BACKEND"] = "vision"
    importlib.reload(app.ocr)
    o = app.ocr
    try:
        check("concurrency is on by default but conservative", 1 < o.CONCURRENCY <= 8)

        slow = {2, 11}

        def fake(image, index, context, **kwargs):
            time.sleep(0.6 if index in slow else 0.05)
            return o.PageText(index, f"page {index}")

        images = [b"x"] * 11
        with mock.patch.object(o, "_read_page_vision", side_effect=fake):
            o.CONCURRENCY = 1
            start = time.time()
            serial = o.read_pages(images, "ctx")
            serial_time = time.time() - start

            o.CONCURRENCY = 4
            start = time.time()
            parallel = o.read_pages(images, "ctx")
            parallel_time = time.time() - start

        check("every page still comes back",
              len(parallel) == 11 and len(serial) == 11)
        check("pages stay in page order however they finish",
              [p.index for p in parallel] == list(range(1, 12)))
        check("the text belongs to the right page",
              all(p.text == f"page {p.index}" for p in parallel))
        check(f"parallel is faster ({parallel_time:.2f}s vs {serial_time:.2f}s)",
              parallel_time < serial_time * 0.75)

        # A page that truly cannot be read must still sink the submission
        # rather than leaving a hole the grader would read as a blank answer.
        def one_bad(image, index, context, **kwargs):
            if index == 5:
                raise o.OCRError("page 5 unreadable")
            return o.PageText(index, "ok")

        o.CONCURRENCY = 4
        with mock.patch.object(o, "_read_page_vision", side_effect=one_bad):
            try:
                o.read_pages(images, "ctx")
                check("one unreadable page still refuses the submission", False)
            except o.OCRError as exc:
                check("one unreadable page still refuses the submission",
                      "page 5" in str(exc))
    finally:
        _restore()


def test_transcription_resumes_after_a_kill():
    """A run killed part way must not start over.

    The review page already tells the user to click Grade again "to pick up
    where it left off". Without a page cache that is false: every retry
    re-reads every page, which on this gateway is minutes per attempt and is
    why a job that keeps dying never finishes.
    """
    import importlib
    from unittest import mock
    import app.ocr

    print("\n-- a killed run resumes from the pages already read --")
    os.environ["OCR_BACKEND"] = "vision"
    importlib.reload(app.ocr)
    o = app.ocr
    try:
        store = {}

        class Cache:
            @staticmethod
            def get(key):
                return store.get(key)

            @staticmethod
            def set(key, text):
                store[key] = text

        read = []

        def reader(image, index, context, **kwargs):
            read.append(index)
            return o.PageText(index, f"text {index}")

        images = [bytes([i]) * 50 for i in range(1, 6)]
        with mock.patch.object(o, "_read_page_vision", side_effect=reader):
            first = o.read_pages(images, "ctx", cache=Cache)
            check("a first run reads every page", sorted(read) == [1, 2, 3, 4, 5])
            check("every page is cached", len(store) == 5)

            read.clear()
            second = o.read_pages(images, "ctx", cache=Cache)
            check("a repeat run reads nothing again", read == [])
            check("the resumed text is identical",
                  [p.text for p in first] == [p.text for p in second])

            # The case that matters: killed after three pages.
            store.clear()
            read.clear()
            for image in images[:3]:
                store[o._cache_key(image)] = "already read"
            resumed = o.read_pages(images, "ctx", cache=Cache)
            check("only the unread pages are read again", sorted(read) == [4, 5])
            check("cached pages keep their place in page order",
                  [p.index for p in resumed] == [1, 2, 3, 4, 5]
                  and resumed[0].text == "already read")

            # Progress must still count cached pages, or the UI goes backwards.
            seen = []
            o.read_pages(images, "ctx", progress=lambda d, t: seen.append(d), cache=Cache)
            check("progress counts cached pages too", seen == [1, 2, 3, 4, 5])

        # Changing the reader must invalidate, not silently serve the old text.
        key_before = o._cache_key(images[0])
        o.VISION_MODEL = "some-other-model"
        check("a different model is a different cache entry",
              o._cache_key(images[0]) != key_before)
    finally:
        _restore()


def test_transcription_prompt():
    print("\n-- the transcription prompt defends against a helpful reader --")
    from app import ocr as o
    prompt = o.TRANSCRIPTION_PROMPT
    flat = " ".join(prompt.split())

    check("it forbids correcting the student's mistakes",
          "DO NOT CORRECT MISTAKES" in prompt)
    check("it says why: corrections become unearned marks",
          "marks the student did not earn" in flat)
    check("it forbids answering or grading",
          "NOT grading it and NOT answering it" in flat)
    check("it asks for illegible text to be marked", "[illegible]" in prompt)
    check("it asks for guesses to be marked", "[?guess]" in prompt)
    check("it keeps crossed-out work visible rather than deciding for the grader",
          "[struck:" in prompt and "the grader decides" in flat)
    # Both of these are failures observed on a real handwritten submission
    # rather than hypotheticals: a struck-and-replaced word came back as
    # "outer inner", and LaTeX leaked into 7 of 11 pages while the other 4 were
    # plain text.
    check("it handles a correction, not just a deletion",
          "[struck: outer] inner" in prompt and "NEVER as \"outer inner\"" in flat)
    check("it forbids LaTeX and Markdown markup",
          "No LaTeX and no Markdown" in prompt)
    check("it requires the same symbol to transcribe the same way on every page",
          "Be CONSISTENT across pages" in prompt)
    check("it keeps the student's notation rather than normalising it",
          'If they wrote "lg", write "lg", not "log"' in flat)
    check("it handles a submission that is a figure rather than a page",
          "SOME SUBMISSIONS ARE A FIGURE, NOT A PAGE OF WRITING" in prompt
          and "the title, the axis labels" in flat)
    check("it will not read values off an axis it cannot measure",
          "Do not estimate precise values off an axis" in flat)
    check("it describes diagrams rather than interpreting them",
          "[diagram:" in prompt and "Do not interpret whether it is correct" in flat)


def test_long_scan_is_refused_not_truncated():
    """A scan over the page limit must be refused, never partly transcribed.

    _render_pdf_pages used to rasterise the first MAX_PAGES and stop. That put
    a partial read BELOW the refuse-on-partial-read check instead of above it:
    a 25-page submission was transcribed to page 20, the note said all 25 pages
    had been read, nothing was marked incomplete, and the grader reported the
    last questions as unattempted. The 19.5/100 failure in a new disguise, and
    a silent one -- so this test counts the pages actually read, not just the
    error.
    """
    print("\n-- a scan over the page limit is refused, not silently cut short --")
    ocr, ex, _patch = _with_backend("vision", TRANSCRIPT)
    try:
        from unittest import mock

        limit = ocr.MAX_PAGES
        read = []

        def counting_reader(image, index, context, **kwargs):
            read.append(index)
            return ocr.PageText(index, TRANSCRIPT)

        with mock.patch.object(ocr, "_read_page_vision", side_effect=counting_reader):
            try:
                ex.extract_document("long.pdf", _scanned_pdf(limit + 5), "Midterm")
                check(f"a {limit + 5}-page scan is refused", False)
            except ex.ExtractionError as exc:
                message = str(exc)
                check(f"a {limit + 5}-page scan is refused", True)
                check("the refusal says how many pages and what the limit is",
                      str(limit + 5) in message and str(limit) in message)
                check("the refusal says it was not graded in part",
                      "in part" in message or "not graded" in message)
                check("the refusal names the setting that would allow it",
                      "OCR_MAX_PAGES" in message)
            check(f"no pages were transcribed before refusing (read {len(read)})",
                  read == [])

            # The limit itself must still be gradable, or the fix has moved the
            # cliff rather than removed it.
            read.clear()
            result = ex.extract_document("exact.pdf", _scanned_pdf(3), "Midterm")
            check("a scan within the limit is still transcribed", len(read) == 3)
            check("every page of it reaches the grader",
                  all(f"page {i}" in result.text for i in (1, 2, 3)))
    finally:
        _restore()


def test_refusals():
    print("\n-- unsupported formats are refused by name --")

    for ext, expected in (("doc", "Save As"), ("png", "turned off on this server"),
                          ("zip", "Archives cannot be read"), ("pages", "Export To"),
                          ("html", "submit the .ipynb file itself"),
                          ("heic", "Images cannot be read")):
        try:
            read(f"x.{ext}", b"content")
            check(f".{ext} is refused", False)
        except extract.ExtractionError as exc:
            check(f".{ext} is refused with actionable advice", expected in str(exc))

    try:
        read("x.xyz", b"content")
        check("an unknown extension names the accepted formats", False)
    except extract.ExtractionError as exc:
        check("an unknown extension names the accepted formats",
              ".pdf, .docx, .txt, .ipynb and .py" in str(exc))
        check("the accepted list is built from what is actually enabled",
              ".png" not in str(exc))

    try:
        read("noextension", b"content")
        check("a file with no extension is refused", False)
    except extract.ExtractionError:
        check("a file with no extension is refused", True)


def test_text_entry():
    print("\n-- Canvas text-entry HTML --")

    check("paragraphs stay separated",
          extract.strip_html("<p>Q1: first</p><p>Q2: second</p>") == "Q1: first\n\nQ2: second")
    check("list items stay separated",
          extract.strip_html("<ul><li>one</li><li>two</li></ul>") == "one\n\ntwo")
    check("line breaks survive", "\n" in extract.strip_html("a<br>b"))
    check("entities are decoded", extract.strip_html("<p>a &lt; b &amp; c</p>") == "a < b & c")
    check("table cells are separated",
          "|" in extract.strip_html("<table><tr><td>Q1</td><td>answer</td></tr></table>"))
    check("script content is dropped",
          "alert" not in extract.strip_html("<script>alert(1)</script><p>real</p>"))
    check("empty html is empty", extract.strip_html("") == "")
    check("None is handled", extract.strip_html(None) == "")


def main():
    test_docx_tables()
    test_docx_reading_order()
    test_docx_other_containers()
    test_docx_incomplete_detection()
    test_docx_header_fallback()
    test_pdf()
    test_txt_encodings()
    test_notebook()
    test_handwriting_off_by_default()
    test_handwriting_transcription()
    test_vision_gateway_shape()
    test_an_empty_reply_is_retried_not_fatal()
    test_runaway_generation_is_bounded_and_detected()
    test_pages_are_transcribed_in_parallel()
    test_transcription_resumes_after_a_kill()
    test_transcription_prompt()
    test_long_scan_is_refused_not_truncated()
    test_refusals()
    test_text_entry()

    print()
    if FAILS:
        print(f"{len(FAILS)} failure(s):")
        for f in FAILS:
            print(f"  - {f}")
        return 1
    print("ALL CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
