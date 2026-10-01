#!/usr/bin/env python3
"""Guards the rule that an approved grade always describes the work it is attached to.

No Canvas, no LLM and no browser -- this drives the storage layer directly
against a temporary SQLite file, so it is safe to run anywhere:

    python3 scripts/test_review_integrity.py

Covers the two ways a review goes stale (the student resubmits, or the rubric
changes), that neither one destroys what a faculty member typed, that an
ordinary re-run leaves approvals alone, that owner-scoping still holds, that
the warning reaches the Excel export, and that a database created by an
earlier version of the app upgrades cleanly.
"""
import os
import sqlite3
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

# Set before the first `app` import: Config reads the environment at import time,
# so importing anything from the package first would bind the real
# instance/autograder.sqlite3 and these tests would write into it.
os.environ["DATABASE_PATH"] = os.path.join(tempfile.mkdtemp(), "import-time.sqlite3")
os.environ.setdefault("FLASK_SECRET_KEY", "test-only")

FAILS = []

OWNER, OTHER = "iu-owner", "iu-other"
COURSE, ASSIGNMENT, STUDENT = "101", "202", "3001"


def check(label, condition):
    print(("PASS  " if condition else "FAIL  ") + label)
    if not condition:
        FAILS.append(label)


def make_app(db_path):
    from flask import Flask

    app = Flask(__name__)
    app.config["DATABASE_PATH"] = db_path
    return app


def test_review_lifecycle():
    from app import excel_export, storage

    db_path = os.path.join(tempfile.mkdtemp(), "review.sqlite3")
    app = make_app(db_path)
    storage.init_db(app)

    def grade(attempt, submitted_at, score, student=STUDENT):
        storage.upsert_graded_submission(
            OWNER, COURSE, ASSIGNMENT, "DSA", student, "Test Student", "f.txt",
            score, 50.0, "ai feedback", None,
            submission_attempt=attempt, submitted_at=submitted_at,
        )

    def row(student=STUDENT):
        rows = storage.list_submissions(OWNER, COURSE, ASSIGNMENT)
        return next(r for r in rows if r["student_id"] == student)

    with app.app_context():
        print("\n-- a resubmission expires the approval --")
        grade(1, "2026-08-31T10:00:00Z", 7.0)
        storage.save_faculty_review(OWNER, row()["id"], 7.0, "Largely incomplete.", approved=True)
        check("approval recorded on attempt 1", row()["approved"] == 1 and row()["review_stale"] == 0)

        grade(2, "2026-08-31T11:00:00Z", 50.0)
        r = row()
        check("resubmission withdraws the approval", r["approved"] == 0)
        check("resubmission flags the row for re-review", r["review_stale"] == 1)
        check("faculty score is NOT destroyed", r["faculty_score"] == 7.0)
        check("faculty feedback is NOT destroyed", r["faculty_feedback"] == "Largely incomplete.")
        check("AI columns still refresh", r["ai_score"] == 50.0)
        check("the new attempt is recorded", r["submission_attempt"] == 2)

        print("\n-- re-grading the same attempt disturbs nothing --")
        storage.save_faculty_review(OWNER, r["id"], 48.0, "Good work.", approved=True)
        check("saving a review clears the stale flag", row()["review_stale"] == 0)
        grade(2, "2026-08-31T11:00:00Z", 50.0)
        r = row()
        check("same attempt keeps the approval", r["approved"] == 1 and r["review_stale"] == 0)
        check("same attempt keeps the faculty score", r["faculty_score"] == 48.0)

        print("\n-- rows nobody reviewed are never flagged --")
        grade(1, "2026-08-31T09:00:00Z", 30.0, student="9999")
        grade(2, "2026-08-31T12:00:00Z", 40.0, student="9999")
        check("unreviewed row is not flagged as stale", row("9999")["review_stale"] == 0)

        print("\n-- changing the rubric expires approvals too --")
        check("first rubric save is not a change",
              storage.save_rubric(OWNER, COURSE, ASSIGNMENT, "DSA", "v1", "r.txt", 50.0, "file") is False)
        check("re-saving identical text is not a change",
              storage.save_rubric(OWNER, COURSE, ASSIGNMENT, "DSA", "v1", "r.txt", 50.0, "file") is False)
        check("replacing with different text IS a change",
              storage.save_rubric(OWNER, COURSE, ASSIGNMENT, "DSA", "v2", "r.txt", 50.0, "file") is True)
        expired = storage.mark_reviews_stale(OWNER, COURSE, ASSIGNMENT)
        r = row()
        check("rubric change expires the approval", r["approved"] == 0 and r["review_stale"] == 1)
        check("rubric change keeps the faculty text", r["faculty_score"] == 48.0)
        check("rubric change skips unreviewed rows", expired == 1 and row("9999")["review_stale"] == 0)

        print("\n-- owner scoping still holds --")
        victim_id = row()["id"]
        check("another user cannot edit this row",
              storage.save_faculty_review(OTHER, victim_id, 50.0, "x", True) is False)
        check("victim row untouched", row()["faculty_score"] == 48.0)
        check("mark_reviews_stale is owner-scoped",
              storage.mark_reviews_stale(OTHER, COURSE, ASSIGNMENT) == 0)

        print("\n-- the export carries the warning --")
        import openpyxl
        buffer = excel_export.build_grades_workbook(
            "DSA", storage.list_submissions(OWNER, COURSE, ASSIGNMENT)
        )
        sheet = openpyxl.load_workbook(buffer).active
        header = [c.value for c in sheet[1]]
        check("export has a Needs Re-Review column", "Needs Re-Review" in header)
        check("export has an Attempt column", "Attempt" in header)
        rows = {r[1]: r for r in sheet.iter_rows(min_row=2, values_only=True)}
        check("stale row exports as Needs Re-Review=Yes",
              rows[STUDENT][header.index("Needs Re-Review")] == "Yes")
        check("stale row exports as Approved=No", rows[STUDENT][header.index("Approved")] == "No")


def test_audit_fixes():
    """Regressions for faults found in the September codebase audit."""
    import os
    from openpyxl import load_workbook
    from app import excel_export

    print("\n-- guards added after the codebase audit --")

    # Feedback opening with "=" was written as a live Excel formula.
    row = {"student_name": "S", "student_id": "1", "ai_score": 8.0, "ai_total": 10.0,
           "faculty_score": None, "ai_feedback": "=1+1 see my note", "faculty_feedback": None,
           "approved": 0, "ai_confidence": "high", "ai_disagreement": None,
           "submission_attempt": 1, "review_stale": 0, "extraction_notes": None}
    sheet = load_workbook(excel_export.build_grades_workbook("A", [row])).active
    cells = [c for r in sheet.iter_rows(min_row=2, max_row=2) for c in r
             if str(c.value or "").startswith("=")]
    check("feedback starting with = is exported as text, not a formula",
          bool(cells) and all(c.data_type == "s" for c in cells))
    check("forcing it to text does not alter the feedback",
          cells[0].value == "=1+1 see my note")

    # A relative database path on App Service sits on ephemeral storage.
    import importlib, sys
    saved = {k: os.environ.get(k) for k in ("DATABASE_PATH", "WEBSITE_INSTANCE_ID", "APP_ENV")}
    try:
        for key in ("DATABASE_PATH", "WEBSITE_INSTANCE_ID"):
            os.environ.pop(key, None)
        os.environ["APP_ENV"] = "development"
        sys.modules.pop("app.config", None)
        check("off Azure the database stays local",
              importlib.import_module("app.config").Config.DATABASE_PATH
              == "instance/autograder.sqlite3")

        os.environ["WEBSITE_INSTANCE_ID"] = "some-instance"
        sys.modules.pop("app.config", None)
        check("on Azure the database defaults under /home, which survives a restart",
              importlib.import_module("app.config").Config.DATABASE_PATH
              == "/home/site/data/autograder.sqlite3")

        os.environ["DATABASE_PATH"] = "/explicit/path.sqlite3"
        sys.modules.pop("app.config", None)
        check("an explicit DATABASE_PATH still wins",
              importlib.import_module("app.config").Config.DATABASE_PATH == "/explicit/path.sqlite3")

        # Production with the published default key means forgeable sessions.
        os.environ.pop("DATABASE_PATH", None)
        os.environ.pop("FLASK_SECRET_KEY", None)
        os.environ["APP_ENV"] = "production"
        sys.modules.pop("app.config", None)
        try:
            importlib.import_module("app.config")
            check("production refuses to start without FLASK_SECRET_KEY", False)
        except RuntimeError as exc:
            check("production refuses to start without FLASK_SECRET_KEY",
                  "FLASK_SECRET_KEY" in str(exc))
    finally:
        for key, value in saved.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        sys.modules.pop("app.config", None)
        importlib.import_module("app.config")

    # An unbounded attachment download can exhaust a 1.75 GB instance.
    from app import canvas_client

    class Oversized:
        headers = {"Content-Length": str(500 * 1024 * 1024)}
        def iter_content(self, _): yield b""
        def close(self): pass

    class Streamed:
        headers = {}
        def __init__(self, size): self.size = size
        def iter_content(self, chunk): 
            sent = 0
            while sent < self.size:
                block = min(chunk, self.size - sent)
                sent += block
                yield b"x" * block
        def close(self): pass

    try:
        canvas_client._read_capped(Oversized())
        check("an oversized attachment is refused on Content-Length", False)
    except canvas_client.CanvasAPIError:
        check("an oversized attachment is refused on Content-Length", True)
    try:
        canvas_client._read_capped(Streamed(canvas_client.MAX_ATTACHMENT_BYTES + 1024))
        check("an attachment that lies about its size is cut off mid-stream", False)
    except canvas_client.CanvasAPIError:
        check("an attachment that lies about its size is cut off mid-stream", True)
    check("an ordinary attachment still downloads whole",
          canvas_client._read_capped(Streamed(1024)) == b"x" * 1024)

    # The background job must be able to get a token without a request session.
    from app import auth_utils
    check("a token can be fetched without a session, for the background job",
          callable(getattr(auth_utils, "access_token_for", None)))


def test_partial_submission_is_refused():
    """A file that failed to read must not become a partial grade.

    Observed for real: a handwritten PDF timed out during transcription while
    the two code files beside it read fine, and the submission came back
    19.5/100 saying problems 3-5 were not attempted -- for work the student had
    done. A grade over unknown scope is worse than no grade.
    """
    from unittest import mock
    from app import extract
    from app.routes.rubric import _extract_submission_text

    print("\n-- an unreadable attachment refuses the whole submission --")

    submission = {"attachments": [
        {"filename": "handwritten.pdf", "url": "https://canvas/1"},
        {"filename": "problem1.py", "url": "https://canvas/2"},
    ]}

    def download(url, token=None):
        return b"x" if url.endswith("1") else b"def my_sort(A):\n    return A\n"

    def extract_document(filename, content, assignment_name=""):
        if filename.endswith(".pdf"):
            raise extract.ExtractionError("transcribing its pages failed: read timed out")
        document = extract.Extracted()
        document.text = content.decode()
        return document

    with mock.patch("app.canvas_client.download_attachment", side_effect=download), \
         mock.patch("app.extract.extract_document", side_effect=extract_document):
        text, filenames, error, notes, _ocr = _extract_submission_text(submission, "tok")

    check("no text is handed to the grader", text is None)
    check("the row carries an error instead of a score", bool(error))
    check("the error names what could not be read", "handwritten.pdf" in error)
    check("the error says the submission was not graded", "not graded" in error)
    check("the readable file is still named for the faculty member",
          "problem1.py" in (filenames or ""))

    # The other direction must still work: warnings are not failures.
    def extract_with_warning(filename, content, assignment_name=""):
        document = extract.Extracted()
        document.text = content.decode()
        document.notes.append("1 image could not be read")
        return document

    with mock.patch("app.canvas_client.download_attachment", side_effect=download), \
         mock.patch("app.extract.extract_document", side_effect=extract_with_warning):
        text, _f, error, notes, _ocr = _extract_submission_text(submission, "tok")
    check("a file read WITH WARNINGS is still graded", text is not None and error is None)
    check("the warning still reaches the faculty member",
          any("image could not be read" in n for n in notes))


def test_healthz_reports_configuration():
    """/healthz must say what is running without saying any secret."""
    import os
    from app import create_app

    print("\n-- /healthz tells an operator what is actually deployed --")
    saved = {k: os.environ.get(k) for k in
             ("OCR_BACKEND", "FLASK_SECRET_KEY", "REALLM_API_KEY", "APP_ENV")}
    try:
        os.environ.update(FLASK_SECRET_KEY="a-real-secret-value",
                          REALLM_API_KEY="sk-a-real-api-key", APP_ENV="development")
        os.environ.pop("OCR_BACKEND", None)
        body = create_app().test_client().get("/healthz")
        payload, raw = body.get_json(), body.get_data(as_text=True)

        check("it still reports liveness", payload["status"] == "ok")
        check("it reports whether image reading is on", payload["ocr_enabled"] is False)
        check("it lists what uploads are accepted right now",
              "png" not in payload["accepted_uploads"])
        check("it reports that the secret key IS set", payload["secret_key_set"] is True)
        check("it never reveals the secret key", "a-real-secret-value" not in raw)
        check("it never reveals the API key", "sk-a-real-api-key" not in raw)
    finally:
        for key, value in saved.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value


def test_long_job_is_not_mistaken_for_a_dead_one():
    """A submission that takes ten minutes must not look like a crash.

    Observed for real: with transcription enabled, one 11-page submission ran
    far longer than the 2-minute staleness window. Progress is counted per
    SUBMISSION, so updated_at never moved, the review page announced the job
    had stopped, and it invited a restart -- which started a second job
    alongside the first, both transcribing the same pages.
    """
    from app import storage
    from app.storage import init_db

    print("\n-- a slow job is alive, not stale --")

    app = make_app(os.path.join(tempfile.mkdtemp(), "slowjob.sqlite3"))
    init_db(app)
    user, course, assignment = "iu-slow", "c1", "a1"
    with app.app_context():
        storage.start_grading_job(user, course, assignment, total=1)
        job = storage.get_grading_job(user, course, assignment)
        check("a job starts alive", not storage.job_is_stale(job))

        # Wind updated_at back past the window without touching progress, which is
        # exactly what a long transcription does.
        db = storage.get_db()
        stale_time = (datetime.now(timezone.utc) - timedelta(seconds=storage.STALE_JOB_SECONDS + 60))
        db.execute("UPDATE grading_jobs SET updated_at = ? WHERE iu_user_id = ?",
                   (stale_time.isoformat(), user))
        db.commit()
        check("silence past the window does read as stale",
              storage.job_is_stale(storage.get_grading_job(user, course, assignment)))

        # The heartbeat is what a healthy long job has and a dead one does not.
        storage.touch_grading_job(user, course, assignment)
        job = storage.get_grading_job(user, course, assignment)
        check("a heartbeat revives it without claiming progress",
              not storage.job_is_stale(job) and job["done"] == 0)
        check("the heartbeat leaves the job running", job["status"] == "running")

        # A finished job must not be kept alive by a late beat.
        storage.finish_grading_job(user, course, assignment, "done")
        storage.touch_grading_job(user, course, assignment)
        check("a heartbeat cannot resurrect a finished job",
              storage.get_grading_job(user, course, assignment)["status"] == "done")

        check("the window is comfortably above the heartbeat interval",
              storage.STALE_JOB_SECONDS >= 4 * 30)


def test_canvas_rubric():
    from app import canvas_rubric
    from app.routes.rubric import _combine_rubric_sources

    print("\n-- reading a rubric off the Canvas assignment payload --")
    assignment = {
        "name": "DSA", "points_possible": 40,
        "rubric_settings": {"title": "Midterm Rubric", "points_possible": 50},
        "rubric": [
            {"id": "_1", "description": "Q1 BST", "long_description": "Definition plus complexity.",
             "points": 10.0,
             "ratings": [{"description": "Exceeds", "points": 10.0, "long_description": "All three parts."},
                         {"description": "No Evidence", "points": 0.0}]},
            {"id": "_2", "description": "Q2 Recursion", "points": 10.0, "ratings": []},
        ],
    }
    parsed = canvas_rubric.from_assignment(assignment)
    check("rubric parsed from the assignment payload", parsed is not None)
    check("rubric_settings total beats assignment points_possible", parsed["points_possible"] == 50)
    check("criteria parsed in order",
          [c["description"] for c in parsed["criteria"]] == ["Q1 BST", "Q2 Recursion"])
    check("points render without a trailing .0", "(10 points)" in parsed["text"])
    check("long_description reaches the prompt", "Definition plus complexity." in parsed["text"])
    check("rating levels reach the prompt", "0 points: No Evidence" in parsed["text"])

    summed = canvas_rubric.from_assignment(
        {"rubric": [{"description": "A", "points": 7.0, "ratings": []},
                    {"description": "B", "points": 3.0, "ratings": []}]}
    )
    check("total falls back to summing the criteria", summed["points_possible"] == 10.0)
    # A Full Marks / No Marks criterion, which is how the EDA assignment's nine
    # steps are scored. Half marks for "removed the duplicate rows" is not a
    # grade a faculty member can defend, so the scale has to reach the prompt.
    steps = canvas_rubric.from_assignment({"rubric": [
        {"description": "Step 1", "long_description": "Open the file and read the dataset.",
         "points": 1.0, "ratings": [{"description": "Full Marks", "points": 1.0},
                                    {"description": "No Marks", "points": 0.0}]},
        {"description": "Step 2", "points": 2.0,
         "ratings": [{"description": "Proficient", "points": 2.0},
                     {"description": "Developing", "points": 1.0},
                     {"description": "No Marks", "points": 0.0}]},
    ]})
    check("a Full Marks / No Marks criterion is recognised as all-or-nothing",
          steps["criteria"][0]["binary"] is True)
    check("a three-level criterion is left as a scale",
          steps["criteria"][1]["binary"] is False)
    check("the all-or-nothing rule reaches the prompt",
          "ALL OR NOTHING: award exactly 1 or exactly 0" in steps["text"])
    check("a graded scale gets no all-or-nothing rule",
          steps["text"].count("ALL OR NOTHING") == 1)
    check("the hidden longer description is what carries the actual requirement",
          "Open the file and read the dataset." in steps["text"])

    check("no rubric attached -> None", canvas_rubric.from_assignment({"name": "x"}) is None)
    check("empty rubric list -> None", canvas_rubric.from_assignment({"rubric": []}) is None)
    check("malformed rubric -> None", canvas_rubric.from_assignment({"rubric": "nope"}) is None)

    print("\n-- combining the Canvas rubric with an uploaded answer key --")
    combined = _combine_rubric_sources(parsed, "Q1 answer: left < node < right.")
    check("combined text labels the Canvas scoring scheme", "SCORING RUBRIC" in combined)
    check("combined text labels the answer key", "ANSWER KEY" in combined)
    check("combined text carries both bodies",
          "Q1 BST" in combined and "left < node < right" in combined)
    check("canvas-only combine omits the notes heading",
          "ANSWER KEY" not in _combine_rubric_sources(parsed, ""))
    check("file-only combine omits the canvas heading",
          "SCORING RUBRIC" not in _combine_rubric_sources(None, "just notes"))


# The schema as it shipped before submissions carried an attempt. A deployed
# instance has a database in exactly this shape, so the upgrade must add the
# columns in place rather than expecting CREATE TABLE to do it.
PRE_ATTEMPT_SCHEMA = """
CREATE TABLE rubrics (
    iu_user_id TEXT NOT NULL, course_id TEXT NOT NULL, assignment_id TEXT NOT NULL,
    assignment_name TEXT, rubric_text TEXT NOT NULL, rubric_filename TEXT,
    canvas_points_possible REAL, uploaded_at TEXT,
    PRIMARY KEY (iu_user_id, course_id, assignment_id));
CREATE TABLE submissions (
    id INTEGER PRIMARY KEY AUTOINCREMENT, iu_user_id TEXT NOT NULL, course_id TEXT NOT NULL,
    assignment_id TEXT NOT NULL, assignment_name TEXT, student_id TEXT NOT NULL,
    student_name TEXT, submission_filename TEXT, ai_score REAL, ai_total REAL,
    ai_feedback TEXT, ai_breakdown_json TEXT, ai_error TEXT, faculty_score REAL,
    faculty_feedback TEXT, approved INTEGER NOT NULL DEFAULT 0, graded_at TEXT,
    UNIQUE (iu_user_id, course_id, assignment_id, student_id));
"""


def test_upgrade_from_existing_database():
    from app import storage

    print("\n-- upgrading a database created by an earlier version --")
    db_path = os.path.join(tempfile.mkdtemp(), "old.sqlite3")
    con = sqlite3.connect(db_path)
    con.executescript(PRE_ATTEMPT_SCHEMA)
    con.execute(
        "INSERT INTO submissions (iu_user_id, course_id, assignment_id, student_id, student_name,"
        " ai_score, ai_total, faculty_score, faculty_feedback, approved)"
        " VALUES ('u', '101', '202', '3001', 'Test Student', 7.0, 50.0, 7.0, 'old review', 1)"
    )
    con.commit()
    con.close()

    app = make_app(db_path)
    storage.init_db(app)  # must not raise against the existing tables

    with app.app_context():
        db = storage.get_db()
        columns = {r["name"] for r in db.execute("PRAGMA table_info(submissions)")}
        check("submission_attempt added", "submission_attempt" in columns)
        check("submitted_at added", "submitted_at" in columns)
        check("review_stale added", "review_stale" in columns)
        check("rubric_source added",
              "rubric_source" in {r["name"] for r in db.execute("PRAGMA table_info(rubrics)")})

        r = storage.list_submissions("u", "101", "202")[0]
        check("the existing row survived", r["faculty_score"] == 7.0 and r["approved"] == 1)
        check("review_stale backfilled to 0", r["review_stale"] == 0)
        check("attempt starts unknown", r["submission_attempt"] is None)

        # The stored attempt is NULL, so the first run after the upgrade cannot
        # prove the work is unchanged. It re-checks rather than assuming.
        storage.upsert_graded_submission(
            "u", "101", "202", "DSA", "3001", "Test Student", "f.txt", 50.0, 50.0, "new", None,
            submission_attempt=2, submitted_at="2026-08-31T11:00:00Z",
        )
        r = storage.list_submissions("u", "101", "202")[0]
        check("a legacy approval is re-checked once after the upgrade",
              r["approved"] == 0 and r["review_stale"] == 1)
        check("legacy faculty text preserved", r["faculty_feedback"] == "old review")

        storage.init_db(app)
        check("init_db is idempotent", True)


def test_stale_row_prefill():
    """A stale row must not offer the old approved score as the default answer."""
    import re

    from app import create_app, storage

    # create_app() reads the import-time path; point it at a file of its own so
    # this test never sees rows left by the ones above.
    app = create_app()
    app.config["DATABASE_PATH"] = os.path.join(tempfile.mkdtemp(), "prefill.sqlite3")
    storage.init_db(app)

    def render(rows):
        return app.jinja_env.get_template("review.html").render(
            course_id="1", assignment_id="2", assignment_name="DSA", submissions=rows,
            total_mismatch=None, stale_reviews=sum(1 for r in rows if r["review_stale"]),
            job=None, job_stale=False, csrf_token=lambda: "x", url_for=lambda *a, **k: "#",
        )

    def rows_for(student=None):
        rows = storage.list_submissions("u", "1", "2")
        if student:
            rows = [r for r in rows if r["student_id"] == student]
        for r in rows:
            r["breakdown"] = []
        return rows

    def score_input(html):
        return re.search(r'name="faculty_score" value="([^"]*)"', html).group(1)

    def feedback_box(html):
        return re.search(r'name="faculty_feedback">(.*?)</textarea>', html, re.S).group(1)

    with app.test_request_context("/"):
        print("\n-- a stale row starts from the new AI result, not the old approval --")
        # Approved 50 on attempt 2, then a wrong-exam file arrives as attempt 3.
        storage.upsert_graded_submission(
            "u", "1", "2", "DSA", "s1", "Test Student", "good.txt", 50.0, 50.0,
            "Excellent work.", None, submission_attempt=2, submitted_at="t2")
        row_id = storage.list_submissions("u", "1", "2")[0]["id"]
        storage.save_faculty_review("u", row_id, 50.0, "Excellent Work", True)
        storage.upsert_graded_submission(
            "u", "1", "2", "DSA", "s1", "Test Student", "rest.txt", 0.0, 50.0,
            "Does not address any question.", None, submission_attempt=3, submitted_at="t3")

        html = render(rows_for())
        check("stale row prefills the new AI score", score_input(html) == "0.0")
        check("stale row prefills the new AI feedback", "Does not address any question." in feedback_box(html))
        check("the old approved score is not in the input", 'name="faculty_score" value="50.0"' not in html)
        check("the old score is still shown read-only", "Earlier review" in html and ">50.0<" in html)
        check("the old comments are still reachable", "Your earlier comments" in html and "Excellent Work" in html)
        check("the approval is unticked", 'name="approved" checked' not in html)
        check("the database still holds the faculty score", rows_for()[0]["faculty_score"] == 50.0)

        print("\n-- an ordinary reviewed row behaves exactly as before --")
        storage.save_faculty_review("u", row_id, 42.0, "Solid answer.", True)
        html = render(rows_for())
        check("normal row prefills the faculty score", score_input(html) == "42.0")
        check("normal row prefills the faculty feedback", "Solid answer." in feedback_box(html))
        check("normal row shows no earlier-review notice", "Earlier review" not in html)
        check("normal row keeps its tick", 'name="approved" checked' in html)

        storage.upsert_graded_submission(
            "u", "1", "2", "DSA", "s2", "Other", "f.txt", 31.0, 50.0, "Partial.", None,
            submission_attempt=1, submitted_at="t1")
        check("unreviewed row prefills the AI score", score_input(render(rows_for("s2"))) == "31.0")


def test_triage_ordering():
    """The queue must lead with the rows where faculty attention changes an outcome."""
    from app.routes.review import _attention_bucket, _sort_key

    print("\n-- the review queue is ordered by what needs attention --")

    def row(name, confidence=None, error=None, stale=0, notes=None, disagreement=None):
        return {"student_name": name, "ai_confidence": confidence, "ai_error": error,
                "review_stale": stale, "extraction_notes": notes,
                "ai_disagreement": disagreement}

    check("a failed row is top priority",
          _attention_bucket(row("a", error="could not read file")) == "error")
    check("a stale review outranks confidence",
          _attention_bucket(row("a", confidence="high", stale=1)) == "stale")
    check("otherwise the row takes the grader's confidence",
          _attention_bucket(row("a", confidence="low")) == "low")
    check("an ungraded confidence defaults to medium",
          _attention_bucket(row("a")) == "medium")
    # A grade computed from a file we only partly read needs a human eye no
    # matter how sure the model said it was.
    check("an extraction warning overrides high confidence",
          _attention_bucket(row("a", confidence="high", notes="2 images unreadable")) == "low")
    # Measured disagreement between two readings beats the model's own claim of
    # certainty -- the claim proved close to worthless in practice.
    check("disagreement between passes overrides high confidence",
          _attention_bucket(row("a", confidence="high", disagreement=3.0)) == "low")
    check("no disagreement leaves the reported confidence alone",
          _attention_bucket(row("a", confidence="high", disagreement=0.0)) == "high")

    rows = [
        row("Eve", confidence="high"),
        row("Bob", confidence="low"),
        row("Dan", confidence="medium"),
        row("Amy", error="download failed"),
        row("Cal", confidence="high", stale=1),
    ]
    ordered = [r["student_name"] for r in sorted(rows, key=_sort_key)]
    check("errors, then stale, then low, then medium, then high",
          ordered == ["Amy", "Cal", "Bob", "Dan", "Eve"])

    same = [row("Zoe", confidence="low"), row("Ann", confidence="low")]
    check("ties fall back to student name",
          [r["student_name"] for r in sorted(same, key=_sort_key)] == ["Ann", "Zoe"])


def test_anchor_pipeline():
    """Approved rows must reach the grader as calibration examples."""
    from app import storage
    from app.routes.rubric import _build_anchor_examples

    print("\n-- approved grades become calibration examples on the next run --")

    db_path = os.path.join(tempfile.mkdtemp(), "anchors.sqlite3")
    app = make_app(db_path)
    storage.init_db(app)

    # Canvas text-entry submissions, so no download is needed.
    submissions = [
        {"user_id": 1, "body": "<p>first answer</p>", "attempt": 1},
        {"user_id": 2, "body": "<p>second answer</p>", "attempt": 1},
        {"user_id": 3, "body": "<p>third answer</p>", "attempt": 1},
    ]

    with app.app_context():
        for student_id in ("1", "2", "3"):
            storage.upsert_graded_submission(
                OWNER, COURSE, ASSIGNMENT, "DSA", student_id, f"Student {student_id}", None,
                20.0, 50.0, "ai feedback", None, submission_attempt=1, submitted_at="t1")

        check("no approvals yet means no anchors",
              _build_anchor_examples(OWNER, COURSE, ASSIGNMENT, submissions, "token") == [])

        rows = {r["student_id"]: r["id"] for r in storage.list_submissions(OWNER, COURSE, ASSIGNMENT)}
        storage.save_faculty_review(OWNER, rows["1"], 11.0, "much too vague", approved=True)
        storage.save_faculty_review(OWNER, rows["3"], 46.0, "thorough", approved=True)
        # Scored but never approved -- must not be used as an exemplar.
        storage.save_faculty_review(OWNER, rows["2"], 30.0, "fine", approved=False)

        anchors = _build_anchor_examples(OWNER, COURSE, ASSIGNMENT, submissions, "token")
        check("only approved rows become anchors", len(anchors) == 2)
        check("the anchors carry the faculty score, not the AI score",
              sorted(a["score"] for a in anchors) == [11.0, 46.0])
        check("the anchors carry the faculty's own comments",
              {a["feedback"] for a in anchors} == {"much too vague", "thorough"})
        check("the graded text is recovered from the Canvas submission",
              {a["text"] for a in anchors} == {"first answer", "third answer"})
        check("an unapproved row is excluded",
              all("second answer" != a["text"] for a in anchors))

        # A stale approval describes work that has been replaced, so it is not a
        # standard to calibrate against.
        storage.upsert_graded_submission(
            OWNER, COURSE, ASSIGNMENT, "DSA", "3", "Student 3", None,
            20.0, 50.0, "ai feedback", None, submission_attempt=2, submitted_at="t2")
        anchors = _build_anchor_examples(OWNER, COURSE, ASSIGNMENT, submissions, "token")
        check("a stale approval is not used as an anchor",
              [a["score"] for a in anchors] == [11.0])

        # A student graded earlier who is missing from this fetch is skipped
        # rather than crashing the run.
        anchors = _build_anchor_examples(OWNER, COURSE, ASSIGNMENT, submissions[1:], "token")
        check("an anchor absent from the current fetch is skipped", anchors == [])


def main():
    test_review_lifecycle()
    test_triage_ordering()
    test_anchor_pipeline()
    test_audit_fixes()
    test_partial_submission_is_refused()
    test_healthz_reports_configuration()
    test_long_job_is_not_mistaken_for_a_dead_one()
    test_canvas_rubric()
    test_upgrade_from_existing_database()
    test_stale_row_prefill()

    print()
    if FAILS:
        print(f"{len(FAILS)} failure(s):")
        for f in FAILS:
            print(f"  - {f}")
        return 1
    print("ALL CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
