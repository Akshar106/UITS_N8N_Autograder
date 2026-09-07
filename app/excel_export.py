import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

COLUMNS = [
    ("Student Name", "student_name"),
    ("Student ID", "student_id"),
    ("AI Score", "ai_score"),
    ("AI Total", "ai_total"),
    ("Final Score", "final_score"),
    ("Percentage", "percentage"),
    ("Approved", "approved_label"),
    ("AI Confidence", "ai_confidence"),
    ("Attempt", "submission_attempt"),
    ("Needs Re-Review", "stale_label"),
    ("Feedback", "final_feedback"),
]


def build_grades_workbook(assignment_name, submissions):
    wb = Workbook()
    ws = wb.active
    ws.title = "Grades"[:31]

    ws.append([label for label, _ in COLUMNS])
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"

    for s in submissions:
        final_score = s["faculty_score"] if s["faculty_score"] is not None else s["ai_score"]
        final_feedback = s["faculty_feedback"] if s["faculty_feedback"] else s["ai_feedback"]
        total = s["ai_total"]
        percentage = round((final_score / total) * 100, 1) if final_score is not None and total else None

        row = {
            "student_name": s["student_name"],
            "student_id": s["student_id"],
            "ai_score": s["ai_score"],
            "ai_total": s["ai_total"],
            "final_score": final_score,
            "percentage": percentage,
            "approved_label": "Yes" if s["approved"] else "No",
            "ai_confidence": s.get("ai_confidence") or "",
            "submission_attempt": s.get("submission_attempt"),
            # Carried into the sheet as well as the web page: whoever transcribes
            # these into Canvas may not be the person who ran the grading, and
            # this is the column that says "do not type this one in yet".
            "stale_label": "Yes" if s.get("review_stale") else "",
            "final_feedback": final_feedback or s.get("ai_error") or "",
        }
        ws.append([row[key] for _, key in COLUMNS])

    for col_idx, (label, _) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        width = 40 if label == "Feedback" else max(14, len(label) + 4)
        ws.column_dimensions[letter].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
